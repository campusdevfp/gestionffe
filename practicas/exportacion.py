"""Exportación de cualquier listado a Excel (.xlsx) o PDF.

Cada informe se define como una lista de cabeceras y una lista de filas ya formateadas,
de modo que ambos formatos salen del mismo sitio.
"""
import io
from datetime import date, datetime

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

TINTA = colors.HexColor("#1d2a30")
PIZARRA = colors.HexColor("#1f5566")
LINEA = colors.HexColor("#d7dfe1")


def _texto(valor):
    if valor is None or valor == "":
        return "—"
    if isinstance(valor, datetime):
        return timezone.localtime(valor).strftime("%d/%m/%Y %H:%M")
    if isinstance(valor, date):
        return valor.strftime("%d/%m/%Y")
    return str(valor)


def a_excel(titulo, cabeceras, filas):
    libro = Workbook()
    hoja = libro.active
    hoja.title = titulo[:31]
    hoja.append(cabeceras)
    for celda in hoja[1]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="1F5566")
        celda.alignment = Alignment(vertical="center")
    for fila in filas:
        hoja.append([_texto(v) if not isinstance(v, (int, float)) else v for v in fila])
    for i, cabecera in enumerate(cabeceras, start=1):
        ancho = max([len(str(cabecera))] + [len(_texto(f[i - 1])) for f in filas] or [10])
        hoja.column_dimensions[get_column_letter(i)].width = min(max(ancho + 2, 12), 45)
    hoja.freeze_panes = "A2"
    hoja.auto_filter.ref = hoja.dimensions
    flujo = io.BytesIO()
    libro.save(flujo)
    return flujo.getvalue()


def a_pdf(titulo, cabeceras, filas, subtitulo=""):
    flujo = io.BytesIO()
    doc = SimpleDocTemplate(flujo, pagesize=landscape(A4), title=titulo,
                            leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    estilos = getSampleStyleSheet()
    estilos["Title"].textColor = TINTA
    estilos["Title"].alignment = 0
    celda = estilos["BodyText"].clone("celda")
    celda.fontSize = 7.5
    celda.leading = 9.5
    cabecera_estilo = celda.clone("cabecera")
    cabecera_estilo.textColor = colors.white
    cabecera_estilo.fontName = "Helvetica-Bold"

    elementos = [Paragraph(titulo, estilos["Title"])]
    pie = f"{subtitulo + ' · ' if subtitulo else ''}{len(filas)} registros · generado el {timezone.localtime():%d/%m/%Y %H:%M}"
    elementos += [Paragraph(pie, estilos["Normal"]), Spacer(1, 6 * mm)]

    datos = [[Paragraph(str(c), cabecera_estilo) for c in cabeceras]]
    datos += [[Paragraph(_texto(v), celda) for v in fila] for fila in filas] or [[Paragraph("Sin datos", celda)] + [""] * (len(cabeceras) - 1)]
    tabla = Table(datos, repeatRows=1, hAlign="LEFT")
    tabla.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), PIZARRA),
        ("GRID", (0, 0), (-1, -1), 0.4, LINEA),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f7f7")]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elementos.append(tabla)
    doc.build(elementos)
    return flujo.getvalue()
